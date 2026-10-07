"""Validated, reversible AoD translation builds. Never modify reference sources."""
from __future__ import annotations
import argparse, csv, hashlib, html, json, os, re, shutil, struct, subprocess, sys, zipfile
import xml.etree.ElementTree as ET
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from dso48 import DSO, decode_text

ROOT=Path(__file__).resolve().parent
INVENTORY=ROOT/'outputs/source_inventory_final_20261004'
SELF=ROOT/'self_service'
LAYOUT_FILE=ROOT/'aod_layout.json'
LAYOUT=json.loads(LAYOUT_FILE.read_text(encoding='utf-8')) if LAYOUT_FILE.exists() else {}
if LAYOUT:
    INVENTORY=(ROOT/LAYOUT['translations']).resolve()
    SELF=(ROOT/LAYOUT['state']).resolve()

def storage(name):
    return (ROOT/LAYOUT[name]).resolve() if LAYOUT else SELF/name

JAPANESE_FONT='BIZ UDPMincho Medium'
csv.field_size_limit(10000000)
def digest(b):return hashlib.sha256(b).hexdigest()
def source_equal(a,b):return a.replace('\r\n','\n')==b.replace('\r\n','\n')
def loadcsv(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def writecsv(p,rows,fields):
    p.parent.mkdir(parents=True,exist_ok=True)
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
def savejson(p,data):p.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
def fail(message):raise ValueError(message)
def safe(root,rel):
    rel=rel.replace('\\','/')
    if not rel or ':' in rel or rel.startswith('/') or '..' in rel.split('/'):fail('Unsafe path: '+rel)
    result=(root/rel).resolve()
    if not result.is_relative_to(root.resolve()):fail('Path escaped root: '+rel)
    return result
def tokens(s):
    # Torque numbered substitutions must not consume the next word ("%1 seconds").
    return re.findall(r'<[^>]+>|%\d+(?![\d$])|%(?:\d+\$)?[-+#0.\d]*[sdfiu]|%[A-Za-z_][A-Za-z_0-9]*|\\[nrt]|\\c\d|\[[^\]\r\n]+\]',s)
def validate_target(source,target,ident):
    if '\x00' in target or '\ufffd' in target:fail('Invalid character: '+ident)
    if tokens(source)!=tokens(target):fail('Protected token/tag mismatch: '+ident)

BLOCKED_STATUSES={'technical_review','do_not_translate','excluded','hold'}

def blocked(row):
    return row.get('status','').strip().lower() in BLOCKED_STATUSES

def select_translations(rows,include_drafts=False):
    selected=[];counts={'skipped_hold_rows':0,'draft_rows_used':0}
    for original in rows:
        if blocked(original):counts['skipped_hold_rows']+=1;continue
        row=dict(original);target=row.get('target_ja','')
        if not target.strip() and include_drafts and row.get('status')=='draft':
            target=row.get('target_draft','')
            if target.strip():counts['draft_rows_used']+=1
        if not target.strip():continue
        validate_target(row['source_raw'],target,row['id'])
        row['target_ja']=target;selected.append(row)
    return selected,counts
def read_table(p,required):
    if p.suffix.lower()=='.csv':rows=loadcsv(p)
    elif p.suffix.lower()=='.xlsx':
        import openpyxl
        w=openpyxl.load_workbook(p,read_only=True,data_only=False);matches=[]
        try:
            for s in w.worksheets:
                it=iter(s.iter_rows(values_only=True));head=next(it,())
                names=[str(v).strip() if v is not None else '' for v in head]
                if required.issubset(names):
                    if len([x for x in names if x])!=len(set(x for x in names if x)):fail('Duplicate Excel headers: '+s.title)
                    values=[];formula_found=False
                    for line in it:
                        if not any(v is not None for v in line):continue
                        row={k:('' if v is None else str(v)) for k,v in zip(names,line) if k}
                        if row.get('target_ja','').startswith('='):formula_found=True
                        if row.get('id','').strip():values.append(row)
                    if formula_found:fail('Translation formula is not allowed: '+s.title)
                    matches.append(values)
        finally:w.close()
        if len(matches)!=1:fail('Excel must have exactly one matching table at row 1: '+str(p))
        rows=matches[0]
    else:fail('Use CSV or XLSX: '+str(p))
    if not rows or not required.issubset(rows[0]):fail('Missing required columns: '+str(p))
    return rows
def merge(baseline,incoming,sourcecol):
    result=[dict(r) for r in baseline];byid={r['id']:r for r in result};seen=set()
    for r in incoming:
        ident=r['id']
        if ident in seen:fail('Duplicate ID: '+ident)
        seen.add(ident)
        if ident not in byid:fail('Unknown ID: '+ident)
        b=byid[ident]
        if not source_equal(r[sourcecol],b[sourcecol]):fail('Source changed: '+ident)
        if r.get('source_hash') and r['source_hash']!=b.get('source_hash'):fail('Source hash changed: '+ident)
        # Blank cells do not erase approved existing translations.
        if r.get('target_ja','').strip():b['target_ja']=r['target_ja']
        # Incoming files may add a hold, but cannot release a canonical hold.
        if blocked(r):b['status']=r['status']
        if not blocked(b) and r.get('target_draft','').strip():
            b['target_draft']=r['target_draft']
            if not b.get('target_ja','').strip() and r.get('status')=='draft':b['status']='draft'
    for r in result:
        if r.get('source_hash'):
            normalized=r[sourcecol].replace('\r\n','\n')
            hashes={digest(r[sourcecol].encode()),digest(normalized.encode()),digest(normalized.replace('\n','\r\n').encode())}
            if r['source_hash'] not in hashes:fail('Invalid source hash: '+r['id'])
        if not blocked(r) and r['target_ja'].strip():validate_target(r[sourcecol],r['target_ja'],r['id'])
    return result
def choose(inputs,stem):
    files=[inputs/(stem+ext) for ext in ['.csv','.xlsx'] if (inputs/(stem+ext)).is_file()]
    if len(files)>1:fail('Keep only CSV or XLSX, not both: '+stem)
    return files[0] if files else None
def source_asset(source,container,path):
    if container.endswith('.aod'):
        with zipfile.ZipFile(safe(source,container)) as z:return z.read(path)
    return safe(source,('' if container=='root' else container+'/')+path).read_bytes()
def runtime_path(container,path):
    if container.endswith('.aod'):
        prefix=container[:-4]
        return prefix+'/'+path
    return ('' if container=='root' else container+'/')+path
def quote(value):
    return '"'+value.replace('\\','\\\\').replace('"','\\"').replace('\n','\\n').replace('\r','\\r').replace('\t','\\t')+'"'
def unescape(value):
    return re.sub(r'\\([\\"\'nrt])',lambda m:{'n':'\n','r':'\r','t':'\t','"':'"',"'":"'",'\\':'\\'}[m[1]],value)
def rewrite_script(text,entries):
    groups=defaultdict(list)
    for e in entries:
        m=re.match(r'(?:text/)?char:(\d+)(?:/(.*))?$',e['locator'])
        if not m:fail('Unsupported text locator: '+e['locator'])
        groups[int(m[1])].append((m[2],e))
    edits=[]
    for start,group in groups.items():
        m=re.match(r'"((?:\\.|[^"\\])*)"',text[start:],re.S)
        if not m:fail('Quote not found at '+str(start))
        raw=unescape(m[1]);direct=[e for child,e in group if not child];nested=[]
        for child,e in group:
            if child:nested.append(dict(e,locator=child))
        if direct:
            if nested:fail('Overlapping parent/child text edits')
            if any(not source_equal(e['source_raw'],raw) for e in direct):fail('Text source mismatch')
            values={e['target_ja'] for e in direct}
            if len(values)!=1:fail('Conflicting text replacements')
            target=values.pop()
        else:target=rewrite_script(raw,nested)
        edits.append((start,start+len(m[0]),quote(target)))
    for start,end,value in sorted(edits,reverse=True):text=text[:start]+value+text[end:]
    return text
def layout_overrides(d,name):
    """Change individual GUI field loads, never shared string-table entries."""
    result={};objects=[];current=None
    for n,i in enumerate(d.instructions):
        if i.op=='CREATE_OBJECT':
            prev=[j for j in d.instructions[max(0,n-5):n] if j.op in {'LOADIMMED_IDENT','LOADIMMED_STR'}]
            current={'class':(prev[0].identifier or prev[0].text) if prev else '',
                     'name':(prev[-1].identifier or prev[-1].text) if len(prev)>1 else '', 'fields':{}}
            objects.append(current)
        elif current and i.op=='SETCURFIELD':
            loads=[j for j in d.instructions[max(0,n-3):n] if j.op=='LOADIMMED_STR']
            if loads:current['fields'][i.identifier]=loads[-1]
    def setfield(o,key,value):
        if key not in o['fields']:fail('Required layout field missing: '+str(o['name'])+'/'+key)
        result[o['fields'][key].ip+1]=value
    effect_names={'cs_'+stat+'_effect' for stat in ['str','dex','con','per','int','cha']}
    effects=[o for o in objects if (o['name'] or '').lower() in effect_names]
    centers=sorted(int(o['fields']['position'].text.split()[0])+int(o['fields']['extent'].text.split()[0])/2 for o in effects)
    # Reserve 12% of the column spacing; all six fields share the same width.
    # This provides room for signed percentages at 125% without touching fonts.
    column_width=round(min(b-a for a,b in zip(centers,centers[1:]))*.88) if len(centers)>1 else None
    for o in objects:
        fields=o['fields'];obj=o['name'] or ''
        if name.startswith('CharScreen_') and obj.lower() in effect_names:
            w,h=map(int,fields['extent'].text.split());x,y=map(int,fields['position'].text.split())
            extra=max(0,(column_width or round(w*2.5))-w)
            setfield(o,'extent',f'{w+extra} {h}');setfield(o,'position',f'{x-extra//2} {y}')
        if name.startswith('maps/Map') and o['class']=='fbMapButtonCtrl':
            w,h=map(int,fields['extent'].text.split())
            setfield(o,'extent',f'{round(w*1.5)} {h}')
        if name=='TutorialOverlay.english.gui.dso':
            extra=130
            if obj in {'TO_Window','TO_Text','TO_MouseCatcher'}:
                w,h=map(int,fields['extent'].text.split());setfield(o,'extent',f'{w} {h+extra}')
            if obj=='TO_Window':
                x,y=map(int,fields['position'].text.split());setfield(o,'position',f'{x} {max(0,y-extra)}')
            if obj in {'TO_btnPrev','TO_btnNext','TO_btnClose','TO_Page'}:
                x,y=map(int,fields['position'].text.split());setfield(o,'position',f'{x} {y+extra}')
    return result

def patch_dso(blob,entries,name=''):
    d=DSO(blob);tables=[bytearray(x) for x in d.tables];code=list(d.code);groups=defaultdict(list)
    # Rendering-only changes: never alter the protected tags in translation data.
    entries=list(entries)
    direct={(e['locator'].split('/')[0]) for e in entries}
    for n,i in enumerate(d.instructions):
        if i.op!='LOADIMMED_STR' or i.text is None:continue
        loc=('global' if i.table==0 else 'function')+':'+str(i.string_offset)
        font_field=any(j.op=='SETCURFIELD' and j.identifier.lower()=='fonttype'
                       for j in d.instructions[n+1:n+4])
        target=JAPANESE_FONT if font_field else render_fonts(i.text)
        if target!=i.text and loc not in direct:
            entries.append({'locator':loc,'source_raw':i.text,'target_ja':target})
            direct.add(loc)
    for e in entries:
        m=re.match(r'(global|function):(\d+)(?:/(.*))?$',e['locator'])
        if not m:fail('Unsupported DSO locator: '+e['locator'])
        groups[(0 if m[1]=='global' else 1,int(m[2]))].append((m[3],e))
    expected={}
    for (table,offset),group in groups.items():
        raw=d.text(table,offset);direct=[e for child,e in group if not child]
        if direct:
            if any(child for child,e in group):fail('Overlapping DSO parent/child edits')
            if any(not source_equal(e['source_raw'],raw) for e in direct):fail('DSO source mismatch')
            values={e['target_ja'] for e in direct}
            if len(values)!=1:fail('Conflicting DSO replacements')
            target=values.pop()
        else:target=rewrite_script(raw,[dict(e,locator=child) for child,e in group])
        target=render_fonts(target)
        newoffset=len(tables[table]);tables[table].extend(target.encode('utf-8')+b'\0')
        ips=[i.ip+1 for i in d.instructions if i.table==table and i.string_offset==offset]
        if not ips:fail('No DSO instruction references translated string')
        for ip in ips:code[ip]=newoffset;expected[ip]=target
    for ip,target in layout_overrides(d,name).items():
        original=next(i for i in d.instructions if i.ip+1==ip)
        newoffset=len(tables[original.table]);tables[original.table].extend(target.encode('utf-8')+b'\0')
        code[ip]=newoffset;expected[ip]=target
    p=4
    for _ in range(2):n=struct.unpack_from('<I',blob,p)[0];p+=4+n
    floatstart=p
    for _ in range(2):n=struct.unpack_from('<I',blob,p)[0];p+=4+n*8
    floatbytes=blob[floatstart:p];n,linecount=struct.unpack_from('<II',blob,p);p+=8
    for _ in range(n):v=blob[p];p+=1+(4 if v==255 else 0)
    trailer=blob[p:]
    built=struct.pack('<I',48)
    for t in tables:built+=struct.pack('<I',len(t))+bytes(t)
    built+=floatbytes+struct.pack('<II',n,linecount)
    built+=b''.join(bytes([v]) if v<255 else b'\xff'+struct.pack('<I',v) for v in code)+trailer
    check=DSO(built)
    assert check.ids==d.ids and check.lines==d.lines and len(check.code)==len(d.code)
    for ip,(a,b) in enumerate(zip(d.code,check.code)):
        if ip not in expected:assert a==b
    for a,b in zip(d.instructions,check.instructions):
        assert a.ip==b.ip and a.op==b.op and a.identifier==b.identifier and a.function==b.function
        if a.string_offset is not None:assert b.text==expected.get(a.ip+1,a.text)
    return built

def render_fonts(text):
    return re.sub(r'(<font:)[^:>]+(:[^>]*>)',lambda m:m[1]+JAPANESE_FONT+m[2],text,flags=re.I)
def patch_xml(blob,entries):
    text=blob.decode('utf-8-sig');tree=ET.fromstring(text);nodes=list(tree.iter())
    opens=[m for m in re.finditer(r'<!--.*?-->|<([A-Za-z_][\w:.-]*)\b[^>]*>',text,re.S) if m[1]]
    if len(opens)!=len(nodes):fail('XML token/element count differs')
    edits=[];seen=set()
    for e in entries:
        m=re.fullmatch(r'xml_element:(\d+):(.+)',e['locator'])
        if not m:fail('Unsupported XML locator')
        index=int(m[1]);node=nodes[index]
        if index in seen:fail('Duplicate XML replacement')
        seen.add(index)
        if node.tag!=m[2] or node.text!=e['source_raw'] or len(node):fail('XML source mismatch')
        start=opens[index].end();end=text.find('<',start)
        if not source_equal(html.unescape(text[start:end]),e['source_raw']):fail('XML text span mismatch: '+e['id'])
        edits.append((start,end,html.escape(e['target_ja'],quote=False)))
    for start,end,value in sorted(edits,reverse=True):text=text[:start]+value+text[end:]
    ET.fromstring(text)
    return (b'\xef\xbb\xbf' if blob.startswith(b'\xef\xbb\xbf') else b'')+text.encode('utf-8')
def build(config,inputs,include_dialogues=False,include_drafts=False):
    SELF.mkdir(parents=True,exist_ok=True)
    # A failed new build must not accidentally install an older successful one.
    savejson(SELF/'latest_build.json',{'status':'BUILD_IN_PROGRESS','build':None})
    source=Path(config['source']).resolve();game=Path(config['game']).resolve()
    rows=loadcsv(INVENTORY/'reviewed_translation_targets.csv')
    p=choose(inputs,'reviewed_translation_targets')
    if p:rows=merge(rows,read_table(p,{'id','source_raw','target_ja'}),'source_raw')
    else:rows=merge(rows,[],'source_raw')
    p=choose(inputs,'chat_fulltext_review_handoff')
    if p:rows=merge(rows,read_table(p,{'id','source_raw','target_ja'}),'source_raw')
    rows,selection_counts=select_translations(rows,include_drafts)
    stamp=datetime.now().strftime('%Y%m%d_%H%M%S_%f');out=storage('builds')/stamp
    out.mkdir(parents=True);payload=out/'payload';payload.mkdir()
    def put(rel,b):
        dest=safe(payload,rel);dest.parent.mkdir(parents=True,exist_ok=True)
        if dest.exists() and dest.read_bytes()!=b:fail('Output path conflict: '+rel)
        dest.write_bytes(b)
    grouped=defaultdict(list); translated=0
    for r in rows:
        if not r['target_ja'].strip():continue
        translated+=1
        for loc in json.loads(r['all_locations']):grouped[(loc['container'],loc['path'])].append(dict(r,locator=loc['locator']))
    # Complete native profile definitions, including inherited fbText profiles.
    # Previously the external Japanese wrapper used an incomplete name whitelist.
    grouped.setdefault(('art/gui.aod','gameProfiles.english.cs.dso'),[])
    grouped.setdefault(('core','art/gui/profiles.cs.dso'),[])
    # SS_Start searches the selected language, then help.cs (not English).
    # Supply the Japanese alias even when its text has not been translated yet.
    grouped.setdefault(('scripts.aod','data/slides/help.english.cs.dso'),[])
    with zipfile.ZipFile(source/'art/gui.aod') as z:
        for name in z.namelist():
            if name.endswith('.english.gui.dso') and (name.startswith('CharScreen_') or name.startswith('maps/Map') or name=='TutorialOverlay.english.gui.dso'):
                grouped.setdefault(('art/gui.aod',name),[])
    for (container,name),entries in sorted(grouped.items()):
        blob=source_asset(source,container,name)
        unique={}
        for e in entries:
            if e['locator'] in unique and unique[e['locator']]['target_ja']!=e['target_ja']:fail('Conflicting locator')
            unique[e['locator']]=e
        entries=list(unique.values())
        if name.endswith('.dso'):built=patch_dso(blob,entries,name)
        elif name.endswith('.xml'):built=patch_xml(blob,entries)
        else:
            text=decode_text(blob).lstrip('\ufeff');built=render_fonts(rewrite_script(text,entries)).encode('utf-8')
        rel=runtime_path(container,name);put(rel,built)
        # Both loader paths are used by the existing mixed-language test setup.
        if '.english.' in rel:put(rel.replace('.english.','.japanese.'),built)
    dialogue_count=0
    if include_dialogues:
        p=choose(inputs,'dialogue_translation')
        if not p:fail('Put dialogue_translation.csv or .xlsx in input first')
        dr=read_table(p,{'id','source','source_hash','source_file','target_ja','status'})
        for r in dr:
            if r.get('target_ja','').strip():validate_target(r['source'],r['target_ja'],r['id'])
            if 'target_draft' in r:r['target_draft']=''
        temp=out/'dialogue_input.csv';writecsv(temp,dr,list(dr[0]))
        result=subprocess.run([sys.executable,str(ROOT/'build_full_dialogues.py'),'--archive',str(source/'scripts.aod'),'--csv',str(temp),'--out-dir',str(out/'dialogues')],capture_output=True,text=True,encoding='utf-8')
        if result.returncode:fail('Dialogue build failed: '+result.stdout+result.stderr)
        for f in (out/'dialogues').glob('*.xml'):
            put('scripts/data/text/dialogues/japanese/'+f.name,f.read_bytes())
            if (game/'mods/JapaneseDisplayTest/dialogues/english'/f.name).exists():put('mods/JapaneseDisplayTest/dialogues/english/'+f.name,f.read_bytes())
            dialogue_count+=1
    manifest={'status':'BUILD_OK_NOT_GAME_TESTED','built_at':stamp,'source':str(source),'game':str(game),'translated_non_dialogue_rows':translated,'dialogue_files':dialogue_count,'scope':'text_only','include_non_dialogue_drafts':include_drafts,**selection_counts,'english_and_japanese_aliases':True,'requires_existing_font_setup':True,'files':[]}
    for f in sorted(payload.rglob('*')):
        if f.is_file():manifest['files'].append({'path':f.relative_to(payload).as_posix(),'sha256':digest(f.read_bytes())})
    savejson(out/'manifest.json',manifest)
    savejson(SELF/'latest_build.json',{'status':'READY','build':str(out)})
    print('Build ready:',out);print('Translated text:',translated,'; files:',len(manifest['files']));print('Skipped holds:',selection_counts['skipped_hold_rows'],'; drafts used:',selection_counts['draft_rows_used']);print('Text-only build. Game files have NOT been changed.')
    return out
def assert_closed(game):
    if os.name=='nt':
        result=subprocess.run(['tasklist','/FO','CSV','/NH'],capture_output=True,text=True,errors='replace')
        if result.returncode:fail('Cannot check game process. Close AoD before proceeding.')
        if re.search(r'"AoD(?:64)?\.exe"',result.stdout,re.I):fail('Close AoD.exe / AoD64.exe first')
def install(config,out,confirm=True):
    game=Path(config['game']).resolve();source=Path(config['source']).resolve();out=out.resolve()
    if game==source or game.is_relative_to(source) or source.is_relative_to(game):fail('Game and read-only source must be separate')
    if not (game/'AoD.exe').is_file() or not (game/'scripts.aod').is_file():fail('Configured folder is not an AoD install')
    assert_closed(game);manifest=json.loads((out/'manifest.json').read_text(encoding='utf-8'))
    if manifest.get('game') and Path(manifest['game']).resolve()!=game:fail('Build belongs to another game folder; rebuild after changing settings')
    # Verify the whole package before changing any game file.
    for r in manifest['files']:
        if digest(safe(out/'payload',r['path']).read_bytes())!=r['sha256']:fail('Build file hash mismatch: '+r['path'])
        safe(game,r['path'])
    if confirm and input('Install into '+str(game)+'? Type YES: ').strip()!='YES':print('Cancelled');return
    backup=storage('backups')/datetime.now().strftime('%Y%m%d_%H%M%S_%f');backup.mkdir(parents=True)
    state={'game':str(game),'build':str(out),'files':[],'state':'installing'}
    # Back up all destinations and loose source companions before the first write.
    planned=[]
    for r in manifest['files']:
        planned.append((r['path'],r['sha256'],'write'))
        if r['path'].endswith('.dso'):
            companion=r['path'][:-4]
            if safe(game,companion).is_file() and companion not in {x['path'] for x in manifest['files']}:planned.append((companion,None,'remove_companion'))
    seen=set()
    for rel,newhash,action in planned:
        if rel in seen:continue
        seen.add(rel);p=safe(game,rel);exists=p.is_file()
        record={'path':rel,'existed':exists,'old_sha256':digest(p.read_bytes()) if exists else None,'installed_sha256':newhash,'action':action}
        if exists:
            copy=safe(backup/'original',rel);copy.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,copy)
        state['files'].append(record)
    savejson(backup/'restore.json',state)
    try:
        for r in state['files']:
            p=safe(game,r['path']);p.parent.mkdir(parents=True,exist_ok=True)
            if r['action']=='write':shutil.copy2(safe(out/'payload',r['path']),p)
            else:p.unlink()
        for r in state['files']:
            if r['action']=='write' and digest(safe(game,r['path']).read_bytes())!=r['installed_sha256']:fail('Installed hash mismatch')
        state['state']='installed';savejson(backup/'restore.json',state)
        savejson(SELF/'latest_backup.json',{'backup':str(backup)})
    except Exception:
        rollback(backup,force=True);raise
    print('Installed. Backup:',backup);print('Start AoD.exe manually, select Japanese, and check the changed screens.')
def rollback(backup,force=False):
    state=json.loads((backup/'restore.json').read_text(encoding='utf-8'));game=Path(state['game']).resolve();assert_closed(game)
    if state.get('state')=='restored':print('Already restored');return
    for r in state['files']:
        p=safe(game,r['path'])
        if r['existed'] and digest(safe(backup/'original',r['path']).read_bytes())!=r['old_sha256']:fail('Backup hash mismatch')
        if not force:
            expected=r['installed_sha256']
            if expected and (not p.is_file() or digest(p.read_bytes())!=expected):fail('File changed after install; do not overwrite: '+r['path'])
            if not expected and p.exists():fail('Companion recreated; review before restoring: '+r['path'])
    for r in state['files']:
        p=safe(game,r['path'])
        if r['existed']:p.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(safe(backup/'original',r['path']),p)
        elif p.exists():p.unlink()
    state['state']='restored';savejson(backup/'restore.json',state);print('Restored previous files. Backup retained:',backup)
def main():
    p=argparse.ArgumentParser();p.add_argument('action',choices=['build','install','restore','menu']);p.add_argument('--config',type=Path,default=SELF/'settings.json');p.add_argument('--input',type=Path,default=SELF/'input');p.add_argument('--build',type=Path);p.add_argument('--dialogues',action='store_true');p.add_argument('--drafts',action='store_true',help='Include non-dialogue draft text; held rows remain excluded');a=p.parse_args()
    # Draft text is opt-in and applies only to non-dialogue strings.
    if LAYOUT and a.input==SELF/'input':a.input=INVENTORY
    config=json.loads(a.config.read_text(encoding='utf-8-sig'))
    for key in ('source','game'):
        value=Path(config[key])
        if not value.is_absolute():config[key]=str((a.config.parent/value).resolve())
    a.input.mkdir(parents=True,exist_ok=True)
    if a.action=='menu':
        print('1: 確定欄の文字列をビルド（ゲームは変更しません）\n2: 最新ビルドを反映（先にバックアップ）\n3: 直前の反映を元に戻す\n4: 会話の翻訳も含めてビルド\n5: 会話以外の下訳も含めてビルド\n0: 終了')
        choice=input('番号: ').strip();a.action={'1':'build','2':'install','3':'restore','4':'build','5':'build'}.get(choice,'exit');a.dialogues=choice=='4';a.drafts=choice=='5'
    if a.action=='build':build(config,a.input,a.dialogues,a.drafts)
    elif a.action=='install':
        if a.build:out=a.build
        else:
            latest=json.loads((SELF/'latest_build.json').read_text(encoding='utf-8'))
            if latest.get('status')!='READY':fail('最新ビルドは成功していません。先に1または4でビルドしてください。')
            out=Path(latest['build'])
        install(config,out)
    elif a.action=='restore':rollback(Path(json.loads((SELF/'latest_backup.json').read_text(encoding='utf-8'))['backup']))
if __name__=='__main__':
    try:main()
    except Exception as e:print('STOP:',e);sys.exit(1)
